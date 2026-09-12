import re
from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def generate_notes_pptx(topic: str, markdown_content: str, output_path: str | Path) -> str:
    """Generate a clean slide presentation (.pptx) from topic notes."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # completely blank

    # Color definitions
    NAVY = RGBColor(26, 54, 93)
    BLUE_ACCENT = RGBColor(43, 108, 176)
    DARK_GRAY = RGBColor(45, 55, 72)
    MUTED_GRAY = RGBColor(113, 128, 150)

    # 1. Slide 1: Title Slide
    slide = prs.slides.add_slide(blank_layout)
    title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(2.5))
    tf = title_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = topic.strip().title()
    p.font.name = "Arial"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = f"Student Lecture & Study Notes • {datetime.now().strftime('%B %Y')}"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.color.rgb = BLUE_ACCENT
    p2.space_before = Pt(14)

    # Parse sections from markdown
    sections = _parse_markdown_into_sections(markdown_content)

    # 2. Add Content Slides
    for section_title, bullets in sections:
        if not bullets and not section_title:
            continue

        slide = prs.slides.add_slide(blank_layout)

        # Header Banner
        header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(1.0))
        htf = header_box.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = section_title
        hp.font.name = "Arial"
        hp.font.size = Pt(28)
        hp.font.bold = True
        hp.font.color.rgb = NAVY

        # Content Box
        body_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.8))
        btf = body_box.text_frame
        btf.word_wrap = True

        first = True
        # Cap bullets per slide to avoid clutter
        for item in bullets[:6]:
            if first:
                bp = btf.paragraphs[0]
                first = False
            else:
                bp = btf.add_paragraph()

            # Clean markdown formatting from text
            clean_item = re.sub(r"\*\*(.+?)\*\*", r"\1", item)
            clean_item = re.sub(r"`(.+?)`", r"\1", clean_item)
            bp.text = f"•  {clean_item}"
            bp.font.name = "Arial"
            bp.font.size = Pt(17)
            bp.font.color.rgb = DARK_GRAY
            bp.space_before = Pt(12)

    prs.save(str(output_path))
    return str(output_path)


def _parse_markdown_into_sections(markdown_content: str) -> list[tuple[str, list[str]]]:
    """Split markdown into (Section Heading, list of bullet/paragraph items)."""
    sections = []
    current_title = "Overview"
    current_items = []

    for line in markdown_content.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("## ") or stripped.startswith("# "):
            if current_items:
                sections.append((current_title, current_items))
                current_items = []
            current_title = stripped.lstrip("# ").strip()
        elif stripped.startswith("### "):
            # Subheading as an item
            sub_title = stripped.lstrip("# ").strip()
            current_items.append(f"**{sub_title}**")
        elif stripped.startswith("- ") or stripped.startswith("* "):
            current_items.append(stripped[2:].strip())
        elif re.match(r"^\d+\.\s", stripped):
            item = re.sub(r"^\d+\.\s+", "", stripped)
            current_items.append(item)
        elif not stripped.startswith(">"):
            if len(stripped) < 200:
                current_items.append(stripped)

    if current_items:
        sections.append((current_title, current_items))

    return sections
